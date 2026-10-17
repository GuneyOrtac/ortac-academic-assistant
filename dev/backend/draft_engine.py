import os
import re
from html.parser import HTMLParser
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt, Inches, RGBColor
import model_factory
from langchain_core.prompts import PromptTemplate
from langchain_core.output_parsers import JsonOutputParser
from pydantic import BaseModel, Field

class DraftSection(BaseModel):
    baslik: str = Field(description="Bölüm başlığı (örn: Giriş, Literatür Özeti, Tartışma)")
    metin: str = Field(description="Bölümün içeriği (akademik bir dille yazılmış paragraf).")
    eksik_atif_var_mi: bool = Field(description="Bu bölümde literatür atıfı eksik olan ve desteklenmesi gereken iddialı cümleler varsa true.")

class DraftResponse(BaseModel):
    makale_basligi: str = Field(description="Makalenin oluşturulan taslak başlığı")
    bolumler: list[DraftSection] = Field(description="Makale bölümleri listesi")

class DocxHTMLParser(HTMLParser):
    def __init__(self, doc):
        super().__init__()
        self.doc = doc
        self.current_p = None
        self.current_heading_level = None
        self.bold = False
        self.italic = False
        self.underline = False
        self.highlight = False

    def handle_starttag(self, tag, attrs):
        attrs_dict = dict(attrs)
        tag = tag.lower()
        if tag in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6'):
            self.current_heading_level = int(tag[1])
            self.current_p = self.doc.add_heading(level=min(self.current_heading_level, 3))
        elif tag in ('p', 'div'):
            if self.current_heading_level is None:
                self.current_p = self.doc.add_paragraph()
        elif tag == 'li':
            self.current_p = self.doc.add_paragraph(style='List Bullet')
        elif tag in ('strong', 'b'):
            self.bold = True
        elif tag in ('em', 'i'):
            self.italic = True
        elif tag == 'u':
            self.underline = True
        elif tag in ('mark',) or 'eksik-atif' in attrs_dict.get('class', ''):
            self.highlight = True
        elif tag == 'br':
            if self.current_p:
                self.current_p.add_run().add_break()

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6'):
            self.current_heading_level = None
            self.current_p = None
        elif tag in ('p', 'div', 'li'):
            self.current_p = None
        elif tag in ('strong', 'b'):
            self.bold = False
        elif tag in ('em', 'i'):
            self.italic = False
        elif tag == 'u':
            self.underline = False
        elif tag in ('mark', 'span'):
            self.highlight = False

    def handle_data(self, data):
        if not data:
            return
        if self.current_p is None:
            if not data.strip():
                return
            self.current_p = self.doc.add_paragraph()
        
        # Eğer metin içinde doğrudan [KAYNAK?] veya [DİKKAT: EKSİK ATIF] varsa vurgulu göster
        run = self.current_p.add_run(data)
        if self.bold:
            run.bold = True
        if self.italic:
            run.italic = True
        if self.underline:
            run.underline = True
        if self.highlight or '[KAYNAK' in data or '[EKSİK ATIF' in data or '[DİKKAT:' in data:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW

def html_to_docx(html_content: str, taslak_yolu: str, default_baslik: str = "Makale Taslağı") -> str:
    """Jodit editoründen veya taslak motorundan gelen HTML içeriğini Word (.docx) belgesine dönüştürür."""
    doc = Document()
    
    # Sayfa kenar boşlukları (Standart akademik tez / makale düzeni: 2.5 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    
    if not html_content or not html_content.strip():
        doc.add_heading(default_baslik, level=1)
        doc.add_paragraph("Bu taslakta henüz içerik bulunmamaktadır.")
    else:
        # HTML etiketlerini parse et
        parser = DocxHTMLParser(doc)
        try:
            parser.feed(html_content)
        except Exception as e:
            # Fallback: düz metin olarak aktar
            clean_text = re.sub(r'<[^>]+>', '', html_content)
            doc.add_heading(default_baslik, level=1)
            for line in clean_text.splitlines():
                if line.strip():
                    p = doc.add_paragraph()
                    run = p.add_run(line)
                    if '[KAYNAK' in line or '[EKSİK' in line:
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                        
    doc.save(taslak_yolu)
    return os.path.abspath(taslak_yolu)

def taslak_metin_uret(konu: str) -> dict:
    """Verilen konuya göre LLM ile akademik taslak üretir ve editöre aktarılacak zengin HTML metnini döndürür."""
    if not model_factory.is_llm_configured(task="drafting"):
        raise ValueError("Makale taslağı oluşturmak için geçerli bir API anahtarı gereklidir. Lütfen 'Ayarlar' sekmesinden API anahtarınızı kaydediniz.")
    
    llm = model_factory.get_llm(task="drafting")
    parser = JsonOutputParser(pydantic_object=DraftResponse)
    
    template = """Sen profesyonel bir akademik makale yazarı ve OpenDraft motorusun. 
    Kullanıcının verdiği konu hakkında detaylı bir akademik makale taslağı oluştur.
    Makale en az 3 mantıksal bölümden (Örn: Giriş, Yöntem / Literatür Değerlendirmesi, Tartışma ve Sonuç) oluşmalıdır.
    Eğer yazılan metnin içinde bilimsel bir gerçeklik iddiası taşıyan ancak atıf yapılmamış cümleler varsa 'eksik_atif_var_mi' değerini true yap.

    Konu: {konu}

    FORMAT KURALLARI:
    {format_instructions}"""
    
    prompt = PromptTemplate(
        template=template,
        input_variables=["konu"],
        partial_variables={"format_instructions": parser.get_format_instructions()},
    )
    
    chain = prompt | llm | parser
    sonuc = chain.invoke({"konu": konu})
    
    # Jodit editörüne aktarılacak HTML yapısını oluştur
    html_parcalari = [f"<h1>{sonuc['makale_basligi']}</h1>"]
    for bolum in sonuc["bolumler"]:
        html_parcalari.append(f"<h2>{bolum['baslik']}</h2>")
        if bolum.get("eksik_atif_var_mi", False):
            html_parcalari.append(
                f"<p>{bolum['metin']} <span class=\"eksik-atif\">[DİKKAT: EKSİK ATIF / KAYNAK GEREKLİ]</span></p>"
            )
        else:
            html_parcalari.append(f"<p>{bolum['metin']}</p>")
            
    html_tam = "".join(html_parcalari)
    
    return {
        "baslik": sonuc["makale_basligi"],
        "html": html_tam,
        "bolum_sayisi": len(sonuc["bolumler"])
    }

def taslak_ve_word_olustur(konu: str, taslak_yolu: str = "makale_taslagi.docx") -> dict:
    """Verilen konuya göre makale taslağı oluşturup Word belgesi olarak kaydeder."""
    uretilen = taslak_metin_uret(konu)
    dosya_yolu = html_to_docx(uretilen["html"], taslak_yolu, default_baslik=uretilen["baslik"])
    
    return {
        "mesaj": "Makale taslağı Word dosyası olarak başarıyla oluşturuldu.",
        "dosya_yolu": dosya_yolu,
        "bolum_sayisi": uretilen["bolum_sayisi"],
        "baslik": uretilen["baslik"],
        "html": uretilen["html"]
    }
